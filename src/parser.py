"""解析层（Parser）：多格式简历解析 + 结构化提取。

支持 PDF（PyMuPDF）与 Word（python-docx），并可回退读取纯文本。
对外主要暴露 `extract_text` 与 `parse_resume`。
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field, asdict
from typing import Dict, List, Optional

# ---- 可选依赖：缺失时给出友好提示，而非直接崩溃 ----
try:
    import fitz  # PyMuPDF
except Exception:  # pragma: no cover
    fitz = None

try:
    import docx  # python-docx
except Exception:  # pragma: no cover
    docx = None


# --------------------------------------------------------------------------- #
# 原始文本抽取
# --------------------------------------------------------------------------- #
def extract_text_from_pdf(data: bytes) -> str:
    """从 PDF 字节流抽取文本（含多页）。"""
    if fitz is None:
        raise RuntimeError("未安装 PyMuPDF，请先 `pip install PyMuPDF`。")
    text_parts: List[str] = []
    with fitz.open(stream=data, filetype="pdf") as doc:
        for page in doc:
            text_parts.append(page.get_text("text"))
    return "\n".join(text_parts)


def extract_text_from_docx(data: bytes) -> str:
    """从 .docx 字节流抽取文本（段落 + 表格）。"""
    if docx is None:
        raise RuntimeError("未安装 python-docx，请先 `pip install python-docx`。")
    document = docx.Document(io.BytesIO(data))
    parts: List[str] = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_text(filename: str, data: bytes) -> str:
    """根据文件扩展名分派到对应解析器。"""
    name = (filename or "").lower()
    if name.endswith(".pdf"):
        return extract_text_from_pdf(data)
    if name.endswith(".docx"):
        return extract_text_from_docx(data)
    if name.endswith(".doc"):
        raise RuntimeError("暂不支持旧版 .doc，请另存为 .docx 或 PDF 后再上传。")
    # 回退：当作纯文本
    try:
        return data.decode("utf-8", errors="ignore")
    except Exception as e:  # pragma: no cover
        raise RuntimeError(f"无法解析该文件：{e}") from e


# --------------------------------------------------------------------------- #
# 结构化提取
# --------------------------------------------------------------------------- #
@dataclass
class ResumeProfile:
    """结构化简历档案。"""

    raw_text: str = ""
    name: str = ""
    email: str = ""
    phone: str = ""
    education: List[str] = field(default_factory=list)
    experience: List[str] = field(default_factory=list)
    projects: List[str] = field(default_factory=list)
    skills: List[str] = field(default_factory=list)

    def to_dict(self) -> Dict:
        return asdict(self)


EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
PHONE_RE = re.compile(r"(?<!\d)(?:\+?86[- ]?)?1[3-9]\d{9}(?!\d)")
# 常见技能词典（用于无 LLM 时的兜底抽取，可按需扩充）
SKILL_VOCAB = [
    "Python", "Java", "Go", "C++", "C#", "JavaScript", "TypeScript", "SQL",
    "PyTorch", "TensorFlow", "Pandas", "NumPy", "Spark", "Hadoop", "Flink",
    "Docker", "Kubernetes", "K8s", "AWS", "GCP", "Azure", "Redis", "MySQL",
    "PostgreSQL", "MongoDB", "Kafka", "Git", "Linux", "FastAPI", "Django",
    "Flask", "Streamlit", "React", "Vue", "Node.js", "Spring", "Maven",
    "机器学习", "深度学习", "自然语言处理", "NLP", "计算机视觉", "CV",
    "大模型", "LLM", "RAG", "Prompt", "微调", "数据挖掘", "数据分析",
    "数据可视化", "推荐系统", "搜索引擎", "分布式", "高并发", "微服务",
    "产品设计", "项目管理", "敏捷开发", "A/B 测试", "用户增长",
]

SECTION_ALIASES = {
    "education": ["教育", "教育背景", "学历", "教育经历", "Education"],
    "experience": ["工作经历", "工作经验", "实习经历", "经历", "Experience", "Work"],
    "projects": ["项目", "项目经历", "项目经验", "Projects"],
    "skills": ["技能", "专业技能", "技术栈", "Skills"],
}


def _looks_like_heading(line: str) -> Optional[str]:
    """判断一行是否像某个板块的标题，返回板块 key。"""
    stripped = line.strip().strip("：:#-•* ")
    if not stripped or len(stripped) > 12:
        return None
    for key, aliases in SECTION_ALIASES.items():
        for alias in aliases:
            if stripped == alias or stripped.startswith(alias):
                return key
    return None


def parse_resume(text: str) -> ResumeProfile:
    """基于规则的轻量结构化提取（无需 LLM，保证离线可用）。

    提取姓名/邮箱/电话，并按标题切分教育、经历、项目、技能等板块。
    """
    profile = ResumeProfile(raw_text=text or "")
    lines = [ln.rstrip() for ln in (text or "").splitlines()]

    # 联系方式
    m = EMAIL_RE.search(text or "")
    if m:
        profile.email = m.group(0)
    m = PHONE_RE.search(text or "")
    if m:
        profile.phone = m.group(0)

    # 姓名：取前 8 行内第一个短且不含数字/邮箱的行
    for ln in lines[:8]:
        cand = ln.strip()
        if 1 < len(cand) <= 12 and not EMAIL_RE.search(cand) and not any(
            ch.isdigit() for ch in cand
        ) and not _looks_like_heading(cand):
            profile.name = cand
            break

    # 板块切分
    current: Optional[str] = None
    buckets: Dict[str, List[str]] = {k: [] for k in SECTION_ALIASES}
    for ln in lines:
        heading = _looks_like_heading(ln)
        if heading:
            current = heading
            continue
        if current and ln.strip():
            buckets[current].append(ln.strip())

    profile.education = buckets["education"][:10]
    profile.experience = buckets["experience"][:20]
    profile.projects = buckets["projects"][:20]

    # 技能：优先取技能板块，其次全文本词典匹配
    if buckets["skills"]:
        joined = "，".join(buckets["skills"])
        profile.skills = [
            s.strip() for s in re.split(r"[，,、/|\n]+", joined) if s.strip()
        ][:40]
    if not profile.skills:
        low = (text or "").lower()
        profile.skills = [s for s in SKILL_VOCAB if s.lower() in low]

    return profile
